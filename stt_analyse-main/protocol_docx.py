"""Протокол в Markdown сохраняется ещё и документом Word."""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

INLINE_MARKUP = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*")
ORDERED_ITEM = re.compile(r"^\d+\.\s+")
BULLET_ITEM = re.compile(r"^[-*]\s+")


def write_protocol_docx(markdown: str, destination: Path) -> None:
    document = Document()
    _configure_document(document)
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if not line.strip():
            index += 1
            continue
        if line.lstrip().startswith("|"):
            index = _add_table(document, lines, index)
            continue
        if line.startswith("### "):
            _add_heading(document, line[4:].strip(), 13, space_before=8)
        elif line.startswith("## "):
            _add_heading(document, line[3:].strip(), 14, space_before=12)
        elif line.startswith("# "):
            _add_heading(document, line[2:].strip(), 16, space_before=0)
        elif BULLET_ITEM.match(line):
            _add_paragraph(document, BULLET_ITEM.sub("", line).strip(), left_indent=Cm(0.75))
        elif ORDERED_ITEM.match(line):
            _add_paragraph(document, line.strip(), left_indent=Cm(0.75))
        else:
            _add_paragraph(document, line.strip(), first_line_indent=Cm(1.25))
        index += 1
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)
    section.footer_distance = Cm(1.0)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    _set_style_font(normal, "Times New Roman")

    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    footer_run = footer.add_run("Протокол встречи")
    _format_run(footer_run, 10, bold=False)


def _set_style_font(style, name: str) -> None:
    style_element = style.element.get_or_add_rPr()
    fonts = style_element.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        style_element.append(fonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attribute), name)


def _format_run(run, size: int, bold: bool = False, italic: bool = False) -> None:
    run.bold = bold
    run.italic = italic
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run._element.get_or_add_rPr().append(fonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attribute), "Times New Roman")


def _add_inline(paragraph, text: str, size: int, force_bold: bool = False) -> None:
    cursor = 0
    for match in INLINE_MARKUP.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            _format_run(run, size, bold=force_bold)
        if match.group(1) is not None:
            run = paragraph.add_run(match.group(1))
            _format_run(run, size, bold=True)
        else:
            run = paragraph.add_run(match.group(2))
            _format_run(run, size, bold=force_bold, italic=True)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        _format_run(run, size, bold=force_bold)
    if cursor == 0 and not text:
        run = paragraph.add_run("")
        _format_run(run, size, bold=force_bold)


def _add_heading(document: Document, text: str, size: int, space_before: int) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.keep_with_next = True
    _add_inline(paragraph, text, size, force_bold=True)


def _add_paragraph(document: Document, text: str, first_line_indent=Cm(0), left_indent=Cm(0)) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = 3
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.first_line_indent = first_line_indent
    paragraph.paragraph_format.left_indent = left_indent
    _add_inline(paragraph, text, 12)


def _add_table(document: Document, lines: list[str], start: int) -> int:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].lstrip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if cells and not all(cell and set(cell) <= set("-: ") for cell in cells):
            rows.append(cells)
        index += 1
    if not rows:
        return index

    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.autofit = True
    for row_index, row in enumerate(rows):
        for column_index in range(width):
            value = row[column_index] if column_index < len(row) else ""
            cell = table.rows[row_index].cells[column_index]
            paragraph = cell.paragraphs[0]
            paragraph.alignment = 0
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            if paragraph.runs:
                paragraph.runs[0].text = ""
            _add_inline(paragraph, value, 11, force_bold=row_index == 0)
            if row_index == 0:
                _shade_cell(cell, "F2F2F2")
    _mark_header_row(table.rows[0])
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(6)
    return index


def _shade_cell(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def _mark_header_row(row) -> None:
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    properties.append(marker)
